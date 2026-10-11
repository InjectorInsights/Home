"""Injector Insights: sign-up site that saves members to Excel and stores certificates."""
import hmac, io, os, re, threading, uuid, zipfile
from datetime import datetime, timezone
from flask import Flask, jsonify, request, send_file, send_from_directory, Response
from openpyxl import Workbook, load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("DATA_DIR", os.path.join(BASE, "data"))
UPLOAD_DIR = os.path.join(DATA_DIR, "certificates")
XLSX = os.path.join(DATA_DIR, "members.xlsx")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
MAX_FILE = 5 * 1024 * 1024
COLS = ["Member ID", "Submitted (UTC)", "Full name", "Email", "Mobile number", "Role", "Country",
        "Opt-in: market research", "Agreed: Terms, Privacy, Compensation", "Agreed: texts and calls",
        "Confirmed 18+", "Licence / registration no.", "Certificate file", "Certificate original name",
        "Verification updated (UTC)"]
ROLES = {"Physician", "Dentist", "Nurse practitioner", "Physician associate", "Registered nurse",
         "Other licensed injector"}
MAGIC = {".pdf": [b"%PDF"], ".png": [b"\x89PNG"], ".jpg": [b"\xff\xd8\xff"], ".jpeg": [b"\xff\xd8\xff"]}

os.makedirs(UPLOAD_DIR, exist_ok=True)
lock = threading.Lock()
app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE + 512 * 1024


def now():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def clean(v, n=200):
    """Trim and cap length."""
    return (v or "").strip()[:n]


def as_text(ws, row):
    """Force every cell to plain text so Excel never evaluates user input as a formula."""
    for cell in ws[row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def open_wb():
    if os.path.exists(XLSX):
        return load_workbook(XLSX)
    wb = Workbook()
    ws = wb.active
    ws.title = "Members"
    ws.append(COLS)
    for i, c in enumerate(COLS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(16, len(c) + 2)
        ws.cell(1, i).font = ws.cell(1, i).font.copy(bold=True)
    ws.freeze_panes = "A2"
    return wb


def need_admin():
    if not ADMIN_PASSWORD:
        return Response("Admin is disabled. Set ADMIN_PASSWORD.", 503)
    a = request.authorization
    if not a or not hmac.compare_digest(a.password or "", ADMIN_PASSWORD):
        return Response("Sign in", 401, {"WWW-Authenticate": 'Basic realm="Injector Insights admin"'})
    return None


@app.after_request
def headers(r):
    r.headers["X-Content-Type-Options"] = "nosniff"
    r.headers["Referrer-Policy"] = "same-origin"
    if request.path.startswith(("/api", "/admin")):
        r.headers["Cache-Control"] = "no-store"
    return r


@app.errorhandler(413)
def too_big(_):
    return jsonify(error="That file is over 5 MB. Choose a smaller PDF, JPG or PNG."), 413


@app.get("/")
def index():
    return send_from_directory(os.path.join(BASE, "static"), "index.html")


@app.get("/healthz")
def health():
    return "ok"


@app.post("/api/register")
def register():
    d = request.get_json(silent=True) or {}
    name, email, mobile = clean(d.get("name")), clean(d.get("email")), clean(d.get("mobile"), 40)
    role, country = clean(d.get("role")), clean(d.get("country"))
    c = d.get("consents") or {}
    if not name or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email) or len(re.sub(r"\D", "", mobile)) < 7:
        return jsonify(error="Please check your name, email and mobile number."), 400
    if role not in ROLES or not country:
        return jsonify(error="Please choose your role and country."), 400
    if not all(c.get(k) is True for k in ("research", "terms", "contact", "age18")):
        return jsonify(error="All four confirmations are required."), 400
    mid = uuid.uuid4().hex
    with lock:
        wb = open_wb()
        ws = wb["Members"]
        ws.append([mid, now(), name, email, mobile, role, country, "Yes", "Yes", "Yes", "Yes", "", "", "", ""])
        as_text(ws, ws.max_row)
        wb.save(XLSX)
    return jsonify(id=mid)


@app.post("/api/verify")
def verify():
    mid = request.form.get("id", "")
    lic = clean(request.form.get("licence"))
    f = request.files.get("certificate")
    if not re.fullmatch(r"[0-9a-f]{32}", mid):
        return jsonify(error="Your session expired. Please start again."), 400
    saved = orig = ""
    if f and f.filename:
        ext = os.path.splitext(f.filename)[1].lower()
        head = f.stream.read(8)
        f.stream.seek(0)
        if ext not in MAGIC or not any(head.startswith(m) for m in MAGIC[ext]):
            return jsonify(error="Upload a PDF, JPG or PNG file."), 400
        saved, orig = mid + ext, clean(os.path.basename(f.filename))
    with lock:
        wb = open_wb()
        ws = wb["Members"]
        row = next((r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == mid), None)
        if row is None:
            return jsonify(error="Your session expired. Please start again."), 404
        if saved:
            f.save(os.path.join(UPLOAD_DIR, saved))
            ws.cell(row, 13, saved)
            ws.cell(row, 14, orig)
        if lic:
            ws.cell(row, 12, lic)
        ws.cell(row, 15, now())
        as_text(ws, row)
        wb.save(XLSX)
    return jsonify(ok=True)


@app.get("/admin/members.xlsx")
def export():
    deny = need_admin()
    if deny:
        return deny
    with lock:
        wb = open_wb()
        buf = io.BytesIO()
        wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="members.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/admin/certificates.zip")
def certs():
    deny = need_admin()
    if deny:
        return deny
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n in sorted(os.listdir(UPLOAD_DIR)):
            z.write(os.path.join(UPLOAD_DIR, n), n)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="certificates.zip", mimetype="application/zip")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 8000)))
